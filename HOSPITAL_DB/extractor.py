import os
import pandas as pd
import re
from docx import Document
from datetime import datetime

def extraer_datos_docx(ruta_archivo):
    try:
        doc = Document(ruta_archivo)
        contenido = []
        for para in doc.paragraphs:
            if para.text.strip():
                contenido.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                texto_fila = " ".join(cell.text.strip() for cell in row.cells)
                if texto_fila.strip():
                    contenido.append(texto_fila)
        return contenido
    except:
        return []

def limpiar_para_csv(texto):
    if not texto: return "N/E"
    # IMPORTANTE: Quitamos comas, saltos de línea y tabuladores para no romper el formato CSV
    res = texto.replace(',', '.').replace('\n', ' ').replace('\t', ' ').replace('\r', ' ').strip()
    return res

def procesar_formato_comas(directorio):
    pacientes = {}
    re_fecha = re.compile(r'(\d{2}/\d{2}/2026)')
    re_ecu = re.compile(r'(?:ECU[:\s]*)(\d{7})|(?<!\d)(\d{7})(?!\d)')
    re_edad = re.compile(r'(\d+\s?[A|M|a|m][\w\s]*)')
    vocabulario_meds = ['MORFINA', 'BUPRENORFINA', 'FENTANYL', 'DEXMEDETO', 'PARACETAMOL', 'KETOROLACO', 'NALBUFINA']

    if not os.path.exists(directorio):
        print(f"Error: No se encontró la ruta {directorio}")
        return []

    for archivo in os.listdir(directorio):
        if archivo.endswith(".docx") and not archivo.startswith("~$"):
            lineas = extraer_datos_docx(os.path.join(directorio, archivo))
            fecha_actual = None
            tipo_atencion = "CONSULTAS"

            for line in lineas:
                m_fecha = re_fecha.search(line)
                if m_fecha:
                    try: fecha_actual = datetime.strptime(m_fecha.group(1), '%d/%m/2026')
                    except: continue
                    continue

                line_up = line.upper()
                if "HOSPITALIZADOS" in line_up: tipo_atencion = "HOSPITALIZADOS"
                elif "INTERCONSULTA" in line_up: tipo_atencion = "INTERCONSULTAS"
                elif "CENSO ALGOLOGÍA" in line_up: tipo_atencion = "ALGO"
                elif "PALIATIVA" in line_up: tipo_atencion = "PALIA"

                m_ecu = re_ecu.search(line)
                if m_ecu and fecha_actual:
                    ecu = m_ecu.group(1) if m_ecu.group(1) else m_ecu.group(2)
                    if ecu not in pacientes:
                        pacientes[ecu] = {
                            'TIPO': tipo_atencion, 'ECU': ecu, 'EDAD': "N/E",
                            'INICIO': fecha_actual, 'FIN': fecha_actual,
                            'TRATAMIENTOS': [], 'DIAG': "", 'MEDS': [],
                            'FALLECE': 0, 'CONTEO': 0
                        }
                    p = pacientes[ecu]
                    p['CONTEO'] += 1
                    if fecha_actual > p['FIN']: p['FIN'] = fecha_actual
                    if fecha_actual < p['INICIO']: p['INICIO'] = fecha_actual
                    
                    if p['EDAD'] == "N/E":
                        m_edad = re_edad.search(line)
                        if m_edad: p['EDAD'] = limpiar_para_csv(m_edad.group(1))
                    
                    # Diagnóstico y tratamiento
                    diag_limpio = limpiar_para_csv(line)
                    p['DIAG'] = diag_limpio
                    p['TRATAMIENTOS'].append(f"[{fecha_actual.strftime('%d/%m')}]: {diag_limpio}")
                    
                    for med in vocabulario_meds:
                        if med in line_up: p['MEDS'].append(med)
                    if any(x in line_up for x in ["FALLECIO", "FALLECIÓ", "DEFUNCION", "EXITUS"]):
                        p['FALLECE'] = 1

    # Crear lista de strings separados por comas
    lineas_csv = ["TIPO DE ATENCION,ECU,EDAD,FECHA INICIO,FECHA FIN,DIAS ATENCION,CONTEO CONSULTAS,DIAGNOSTICO,TRATAMIENDO,MEDICAMENTO,VECES DADO MEDICAMENTO,FALLECE"]
    
    for ecu, p in pacientes.items():
        campos = [
            p['TIPO'],
            p['ECU'],
            p['EDAD'],
            p['INICIO'].strftime('%d/%m/%Y'),
            p['FIN'].strftime('%d/%m/%Y'),
            str((p['FIN'] - p['INICIO']).days + 1),
            str(p['CONTEO']),
            p['DIAG'],
            " // ".join(p['TRATAMIENTOS']),
            " ".join(set(p['MEDS'])),
            str(len(p['MEDS'])),
            str(p['FALLECE'])
        ]
        lineas_csv.append(",".join(campos))
    
    return lineas_csv

# --- EJECUCIÓN ---
ruta_hgm = r'****************************************'
resultado = procesar_formato_comas(ruta_hgm)

if resultado:
    print("\n" + " TEXTO PARA EXCEL (DELIMITADO POR COMAS) ".center(60, "="))
    for r in resultado:
        print(r)
    print("="*60 + "\n")
